import zipfile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from app.services.parsing.inference import dedupe_header, infer_column_type
from app.services.parsing.jsonsafe import json_safe_row
from app.services.parsing.types import CorruptFile, EmptyWorkbook, ParsedColumn, ParsedSheet, ParseResult

_INFERENCE_SAMPLE_SIZE = 200


def _parse_sheet(ws, sheet_index: int, max_rows: int, sample_row_count: int) -> ParsedSheet:
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if header_row is None or all(v is None for v in header_row):
        return ParsedSheet(
            name=ws.title,
            sheet_index=sheet_index,
            row_count=0,
            columns=[],
            sample_rows=[],
            warning="Sheet is empty",
        )

    column_names = dedupe_header(list(header_row))
    column_values: list[list[object]] = [[] for _ in column_names]
    rows: list[dict] = []
    row_count = 0
    truncated = False

    for row in rows_iter:
        row_count += 1
        if row_count > max_rows:
            truncated = True
            continue

        padded = list(row) + [None] * (len(column_names) - len(row))
        rows.append(json_safe_row(dict(zip(column_names, padded))))
        if row_count <= _INFERENCE_SAMPLE_SIZE:
            for i in range(len(column_names)):
                column_values[i].append(padded[i] if i < len(padded) else None)

    columns = []
    for position, name in enumerate(column_names):
        inferred_type, confidence = infer_column_type(column_values[position])
        columns.append(
            ParsedColumn(name=name, position=position, inferred_type=inferred_type, confidence=confidence)
        )

    warning = f"Row count exceeds {max_rows}; data beyond the limit was not processed" if truncated else None

    return ParsedSheet(
        name=ws.title,
        sheet_index=sheet_index,
        row_count=min(row_count, max_rows),
        columns=columns,
        sample_rows=rows[:sample_row_count],
        rows=rows,
        warning=warning,
    )


def parse_excel(path: str, max_rows: int, sample_row_count: int) -> ParseResult:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, zipfile.BadZipFile) as exc:
        raise CorruptFile(str(exc)) from exc

    try:
        sheet_names = wb.sheetnames
        if not sheet_names:
            raise EmptyWorkbook("Workbook has no worksheets")

        sheets = [
            _parse_sheet(wb[name], idx, max_rows, sample_row_count) for idx, name in enumerate(sheet_names)
        ]
        return ParseResult(sheets=sheets)
    finally:
        wb.close()
