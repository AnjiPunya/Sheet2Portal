import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.db import Base, get_db
from app.main import app


@pytest.fixture(autouse=True)
def _no_real_provider_keys(monkeypatch):
    """Tests must never see (or spend) real API keys from a developer's .env."""
    from app.core.config import settings

    for field in ("groq_api_key", "gemini_api_key", "gemini_api_key_2", "openrouter_api_key", "cerebras_api_key"):
        monkeypatch.setattr(settings, field, "")
    monkeypatch.setattr(settings, "chatbot_provider", "groq")
    monkeypatch.setattr(settings, "chatbot_model", "")


@pytest.fixture(autouse=True)
def _fast_password_hashing(monkeypatch):
    """Cheaper scrypt in tests (the cost is stored in each hash, so verifying
    still works). Also resets the in-memory login throttle between tests."""
    from app.api.routes import auth as auth_route
    from app.core import security

    monkeypatch.setattr(security, "_SCRYPT_N", 2**8)
    auth_route._failures.clear()


@pytest.fixture
def client(tmp_path, monkeypatch):
    """A TestClient wired to an isolated in-memory SQLite DB and a temp upload dir.

    SQLite stands in for Postgres in tests because the models use SQLAlchemy's
    dialect-agnostic Uuid/JSON types -- no Postgres-only features are used.
    """
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    TestSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)

    def override_get_db():
        db = TestSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db

    from app.api.routes import portal as portal_route
    from app.api.routes import workbooks as workbooks_route
    from app.services.storage.local import LocalStorageBackend

    test_storage = LocalStorageBackend(str(tmp_path / "uploads"))
    monkeypatch.setattr(workbooks_route, "storage", test_storage)
    # portal.py did `from workbooks import storage`, which bound its own
    # independent name at import time -- patching workbooks_route.storage
    # above does not cascade to it, so it needs patching separately.
    monkeypatch.setattr(portal_route, "storage", test_storage)

    test_client = TestClient(app)
    test_client.test_sessionmaker = TestSessionLocal
    with test_client:
        # Every API route requires a signed-in user; the default client is
        # the owner of its own tenant. Use `make_client` for other users.
        _sign_up(test_client, DEFAULT_EMAIL)
        yield test_client

    app.dependency_overrides.clear()


DEFAULT_EMAIL = "owner@example.com"
TEST_PASSWORD = "correct-horse-1"


def _sign_up(test_client, email, password=TEST_PASSWORD):
    response = test_client.post("/api/auth/signup", json={"email": email, "password": password})
    assert response.status_code == 201, response.text
    return response.json()


@pytest.fixture
def make_client(client):
    """Factory for extra signed-in clients (each its own user and tenant),
    sharing the same test database. ``role`` changes the new user's role."""
    from app.models.user import User

    opened = []

    def make(email: str, role: str | None = None) -> TestClient:
        other = TestClient(app)
        other.__enter__()
        opened.append(other)
        _sign_up(other, email)
        if role:
            with client.test_sessionmaker() as db:
                db.query(User).filter(User.email == email).update({"role": role})
                db.commit()
        return other

    yield make
    for other in opened:
        other.__exit__(None, None, None)


@pytest.fixture
def anon_client(client):
    """A client with no session, against the same test database."""
    other = TestClient(app)
    with other:
        yield other


@pytest.fixture
def multi_sheet_xlsx(tmp_path):
    path = tmp_path / "workbook.xlsx"
    wb = openpyxl.Workbook()

    projects = wb.active
    projects.title = "Projects"
    projects.append(["Project", "Budget", "Spent", "Start Date"])
    projects.append(["Website Revamp", 10000, 8000, "2024-01-15"])
    projects.append(["Mobile App", 25000, 27000, "2024-02-01"])

    tasks = wb.create_sheet("Tasks")
    tasks.append(["Task", "Project", "Done"])
    tasks.append(["Design mockups", "Website Revamp", True])
    tasks.append(["Build API", "Mobile App", False])

    wb.save(path)
    return str(path)


@pytest.fixture
def empty_sheet_xlsx(tmp_path):
    path = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Blank"
    wb.save(path)
    return str(path)


@pytest.fixture
def duplicate_header_xlsx(tmp_path):
    path = tmp_path / "dupes.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Name", None])
    ws.append(["Alice", "Smith", "extra"])
    wb.save(path)
    return str(path)


@pytest.fixture
def simple_csv(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("Name,Age,Active\nAlice,30,true\nBob,25,false\n")
    return str(path)


@pytest.fixture
def empty_csv(tmp_path):
    path = tmp_path / "empty.csv"
    path.write_text("")
    return str(path)
