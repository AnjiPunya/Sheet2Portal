"""Generates a synthetic retail/e-commerce workbook -- a structurally and
domain-wise different dataset from the construction demo, used to verify the
portal engine is genuinely data-driven rather than tuned to one domain.

Usage: python scripts/generate_demo_data_retail.py [output_path]
"""

import sys
from datetime import date, timedelta

import openpyxl

TODAY = date.today()


def _rel(days: int) -> date:
    return TODAY + timedelta(days=days)


CUSTOMERS = [
    ("Aarav Mehta", "aarav.mehta@example.invalid", "Mumbai", _rel(-400), "Gold", 48200),
    ("Sara Thomas", "sara.thomas@example.invalid", "Bengaluru", _rel(-220), "Silver", 15600),
    ("Imran Qureshi", "imran.q@example.invalid", "Delhi", _rel(-90), "Bronze", 3200),
    ("Nisha Patel", "nisha.patel@example.invalid", "Pune", _rel(-610), "Gold", 61500),
]

PRODUCTS = [
    ("Wireless Mouse", "Accessories", 799, 8, 20),
    ("Mechanical Keyboard", "Accessories", 3499, 3, 10),
    ("27-inch Monitor", "Displays", 15999, 12, 5),
    ("USB-C Hub", "Accessories", 1299, 1, 15),
    ("Laptop Stand", "Accessories", 1899, 25, 10),
    ("Webcam 1080p", "Accessories", 2199, 4, 12),
]

ORDERS = [
    ("ORD-1001", "Aarav Mehta", "Mechanical Keyboard", 1, _rel(-10), "Delivered", 3499),
    ("ORD-1002", "Sara Thomas", "27-inch Monitor", 2, _rel(-5), "Shipped", 31998),
    ("ORD-1003", "Imran Qureshi", "USB-C Hub", 3, _rel(-2), "Pending", 3897),
    ("ORD-1004", "Nisha Patel", "Wireless Mouse", 1, _rel(-1), "Pending", 799),
    ("ORD-1005", "Aarav Mehta", "Webcam 1080p", 2, _rel(-15), "Cancelled", 4398),
    ("ORD-1006", "Sara Thomas", "Laptop Stand", 1, _rel(-8), "Delivered", 1899),
]

REVIEWS = [
    ("Mechanical Keyboard", "Aarav Mehta", 5, "Great tactile feel.", _rel(-8)),
    ("27-inch Monitor", "Sara Thomas", 4, "Good colors, bit heavy.", _rel(-3)),
    ("Laptop Stand", "Sara Thomas", 5, "Sturdy and compact.", _rel(-6)),
]


def build_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()

    customers = wb.active
    customers.title = "Customers"
    customers.append(["Customer Name", "Email", "City", "Signup Date", "Loyalty Tier", "Total Spent"])
    for row in CUSTOMERS:
        customers.append(list(row))

    products = wb.create_sheet("Products")
    products.append(["Product Name", "Category", "Price", "Stock Qty", "Reorder Point"])
    for row in PRODUCTS:
        products.append(list(row))

    orders = wb.create_sheet("Orders")
    orders.append(["Order ID", "Customer Name", "Product Name", "Quantity", "Order Date", "Status", "Total Amount"])
    for row in ORDERS:
        orders.append(list(row))

    reviews = wb.create_sheet("Reviews")
    reviews.append(["Product Name", "Customer Name", "Rating", "Comment", "Review Date"])
    for row in REVIEWS:
        reviews.append(list(row))

    return wb


if __name__ == "__main__":
    output_path = sys.argv[1] if len(sys.argv) > 1 else "demo_retail_portal.xlsx"
    build_workbook().save(output_path)
    print(f"saved {output_path}")
