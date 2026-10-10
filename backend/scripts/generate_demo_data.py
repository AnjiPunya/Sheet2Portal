"""Generates a synthetic construction-management workbook for demoing/testing
the portal generator. Not customer data -- purely for development.

Usage: python scripts/generate_demo_data.py [output_path]
"""

import random
import sys
from datetime import date, timedelta

import openpyxl

random.seed(42)

PROJECTS = [
    ("Green Meadows Apartments", "Residential High-Rise", "Rajesh Sharma", "Whitefield, Bengaluru", "In Progress", 49700000, 24400000, date(2024, 1, 10), date(2025, 6, 30)),
    ("Lakeview Luxury Villas", "Gated Community Villas", "Priya Nair", "Sarjapur Road, Bengaluru", "At Risk", 32000000, 29800000, date(2024, 3, 1), date(2025, 3, 31)),
    ("Orchid Commercial Fit-out", "Commercial Office", "Amit Verma", "Outer Ring Road, Bengaluru", "In Progress", 18500000, 6200000, date(2024, 5, 15), date(2024, 12, 15)),
]

MATERIALS = [
    ("Cement (OPC 53 Grade)", "Bulk Material", "Bags", 120, 500, 420, "Green Meadows Apartments"),
    ("TMT Steel Bars 12mm", "Steel", "Tonnes", 8, 25, 42000, "Green Meadows Apartments"),
    ("Red Clay Bricks", "Masonry", "Units", 62000, 50000, 8, "Lakeview Luxury Villas"),
    ("Ready Mix Concrete M25", "Bulk Material", "Cu.m", 30, 100, 5800, "Lakeview Luxury Villas"),
    ("Electrical Conduit Pipes", "Electrical", "Metres", 1400, 1000, 45, "Orchid Commercial Fit-out"),
    ("Vitrified Floor Tiles", "Finishing", "Sq.ft", 1800, 2000, 85, "Orchid Commercial Fit-out"),
    ("Waterproofing Membrane", "Finishing", "Rolls", 4, 20, 3200, "Green Meadows Apartments"),
]

PURCHASE_REQUESTS = [
    ("PR-1001", "TMT Steel Bars 12mm", "Green Meadows Apartments", "Suresh Kumar", 20, "Pending Approval", 840000, date(2025, 9, 20)),
    ("PR-1002", "Ready Mix Concrete M25", "Lakeview Luxury Villas", "Priya Nair", 60, "Pending Approval", 348000, date(2025, 9, 22)),
    ("PR-1003", "Vitrified Floor Tiles", "Orchid Commercial Fit-out", "Amit Verma", 500, "Approved", 42500, date(2025, 9, 10)),
    ("PR-1004", "Waterproofing Membrane", "Green Meadows Apartments", "Suresh Kumar", 10, "Pending Approval", 32000, date(2025, 9, 25)),
    ("PR-1005", "Electrical Conduit Pipes", "Orchid Commercial Fit-out", "Amit Verma", 300, "Rejected", 13500, date(2025, 9, 5)),
    ("PR-1006", "Red Clay Bricks", "Lakeview Luxury Villas", "Priya Nair", 10000, "Pending Approval", 80000, date(2025, 9, 28)),
]

TODAY = date.today()


def _rel(days: int) -> date:
    """A date `days` from today (negative = past, for realistic overdue/paid examples)."""
    return TODAY + timedelta(days=days)


TASKS = [
    ("Foundation Pouring - Tower A", "Green Meadows Apartments", "Suresh Kumar", "Done", _rel(-90), "High"),
    ("Structural Steel Fixing", "Green Meadows Apartments", "Suresh Kumar", "In Progress", _rel(14), "High"),
    ("MEP Rough-in Villa 4-8", "Lakeview Luxury Villas", "Priya Nair", "Open", _rel(-6), "High"),
    ("Facade Glazing Install", "Orchid Commercial Fit-out", "Amit Verma", "In Progress", _rel(21), "Medium"),
    ("Interior Fit-out Floor 5", "Orchid Commercial Fit-out", "Amit Verma", "Open", _rel(5), "Medium"),
    ("Landscaping Phase 1", "Lakeview Luxury Villas", "Priya Nair", "Open", _rel(30), "Low"),
    ("Final Plumbing Inspection", "Green Meadows Apartments", "Suresh Kumar", "Open", _rel(-2), "High"),
]

EXPENSES = [
    ("Green Meadows Apartments", "Labor", 1850000, _rel(-20), "Paid"),
    ("Green Meadows Apartments", "Equipment Rental", 420000, _rel(-12), "Paid"),
    ("Lakeview Luxury Villas", "Material Procurement", 2100000, _rel(-10), "Paid"),
    ("Lakeview Luxury Villas", "Labor", 980000, _rel(-4), "Pending"),
    ("Orchid Commercial Fit-out", "Design Consultancy", 350000, _rel(-18), "Paid"),
    ("Orchid Commercial Fit-out", "Material Procurement", 610000, _rel(-1), "Pending"),
]

ISSUES = [
    ("Lakeview Luxury Villas", "Soil settlement detected near Block C", "Critical", "Open", _rel(-8)),
    ("Lakeview Luxury Villas", "Vendor delay on precast panels", "High", "Open", _rel(-2)),
    ("Green Meadows Apartments", "Minor crack in parking ramp", "Medium", "Resolved", _rel(-25)),
    ("Orchid Commercial Fit-out", "HVAC duct clash with structural beam", "High", "Open", _rel(-1)),
]


def build_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()

    projects = wb.active
    projects.title = "Projects"
    projects.append(["Project Name", "Type", "Manager", "Location", "Status", "Budget", "Spent", "Start Date", "End Date"])
    for row in PROJECTS:
        projects.append(list(row))

    materials = wb.create_sheet("Materials")
    materials.append(["Material Name", "Category", "Unit", "Quantity On Hand", "Reorder Level", "Unit Cost", "Project"])
    for row in MATERIALS:
        materials.append(list(row))

    prs = wb.create_sheet("Purchase Requests")
    prs.append(["PR ID", "Item", "Project", "Requested By", "Quantity", "Status", "Amount", "Requested Date"])
    for row in PURCHASE_REQUESTS:
        prs.append(list(row))

    tasks = wb.create_sheet("Tasks")
    tasks.append(["Task Name", "Project", "Assigned To", "Status", "Due Date", "Priority"])
    for row in TASKS:
        tasks.append(list(row))

    expenses = wb.create_sheet("Expenses")
    expenses.append(["Project", "Category", "Amount", "Date", "Status"])
    for row in EXPENSES:
        expenses.append(list(row))

    issues = wb.create_sheet("Issues")
    issues.append(["Project", "Title", "Severity", "Status", "Reported Date"])
    for row in ISSUES:
        issues.append(list(row))

    return wb


if __name__ == "__main__":
    output_path = sys.argv[1] if len(sys.argv) > 1 else "demo_construction_portal.xlsx"
    build_workbook().save(output_path)
    print(f"saved {output_path}")
